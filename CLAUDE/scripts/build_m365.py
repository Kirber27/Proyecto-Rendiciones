import json, glob, unicodedata, datetime as dt, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]  # uso: python build_m365.py salida.xlsx  (ejecutar en una carpeta con lists.json, codes.json y seed/b*.json = datos/historial_2026)
F = "Arial"
H_IN = PatternFill("solid", fgColor="1F3864"); H_FX = PatternFill("solid", fgColor="7F7F7F"); H_FL = PatternFill("solid", fgColor="2E75B6")
WH = Font(name=F, bold=True, color="FFFFFF", size=10); BASE = Font(name=F, size=10); BLUE = Font(name=F, size=10, color="0000FF")
CLP = '#,##0;-#,##0;"-"'; DATE = "dd-mm-yyyy"
red = PatternFill("solid", fgColor="F8CBAD"); amber = PatternFill("solid", fgColor="FFE699"); green = PatternFill("solid", fgColor="C6EFCE")

LI = json.load(open("lists.json")); CODES = json.load(open("codes.json"))
tr = lambda s: " ".join(str(s).split()) if s is not None else ""
norm = lambda s: " ".join(unicodedata.normalize("NFD", str(s)).encode("ascii", "ignore").decode().lower().split())
SOL = [tr(x) for x in LI["sol"]]; SUB = [tr(x) for x in LI["sub"]]; CUE = [tr(x) for x in LI["cue"]]; CC = [tr(x) for x in LI["cc"]]; BAN = [tr(x) for x in LI["pag"]]
SUB2 = {"Artículos de Aseo": "Oficina", "Artículos de Oficina": "Oficina", "Bencina": "Operacionales", "Bodega": "Operacionales", "Café / Doctores Colación": "Operacionales", "Comidas y Bebidas": "Operacionales", "DIN Importaciones": "Importaciones / Export.", "Envíos y encomiendas": "Operacionales", "Equipos Electronicos y de Computacion": "Activo fijo", "Equipos de Protección personal": "Oficina", "Estacionamiento": "Operacionales", "Gastos Auto": "Operacionales", "Gastos Furgón": "Operacionales", "Gastos Importacion": "Importaciones / Export.", "Gastos Legales": "Operacionales", "Gastos Otros Países - Chile": "Extras", "Gastos Varios": "Oficina", "Gastos de Vehiculo": "Operacionales", "Gastos de Viajes": "Operacionales", "Honorarios": "Operacionales", "Hospedaje": "Viajes", "Instituto Salud Publica (CDA)": "Importaciones / Export.", "Insumos Dentales": "Otros", "Insumos para Presentaciones": "Operacionales", "Maletas Pelican": "Operacionales", "Miscelaneos Administracion": "Oficina", "Notario y Legal": "Extras", "Otros": "Operacionales", "Otros Gastos de Oficina": "Oficina", "Pasajes Bus - Aéreos": "Operacionales", "Patente": "Operacionales", "Peajes / Estacionamientos": "Operacionales", "Publicidad / Marketing": "Extras", "Pólizas de garantías Licitaciones": "Extras", "Reparaciones y remodelaciones": "Operacionales", "Taxi - Uber": "Operacionales"}
EMP = [("Heca Global SPA", "76723245-4"), ("Hecaglobal Instruments SPA", "76886693-7"), ("Siga Sonriendo SPA", "77422425-4"), ("Sonrie Pro Financial SPA", "77219167-7")]

# ---------------- history
W = []
for f in sorted(glob.glob("seed/b*.json")): W += json.load(open(f))
rends = [dict(id=w["doc_id"], **w["data"]) for w in W if w["collection"] == "rendiciones"]
pagos = [w["data"] for w in W if w["collection"] == "pagos"]
D = lambda s: dt.datetime.strptime(s, "%Y-%m-%d") if s else None
rends.sort(key=lambda r: (r["fecha"], r["id"]))
R_rows, D_rows, P_rows = [], [], []
for r in rends:
    R_rows.append({"ID": r["id"], "Nro": r["nro"], "Solicitante": tr(r["solicitante"]), "Empresa": r["empresa"], "Fecha": D(r["fecha"]),
                   "Estado": "Aprobada", "Origen": "Histórico Excel v1.1"})
    for i, l in enumerate(r["lineas"], 1):
        D_rows.append({"ID": r["id"], "Línea": i, "Fecha": D(l["fecha"]), "Tipo doc": l["tipo"], "Neto": l["neto"] or None, "IVA": round(l["iva"], 2) or None,
                       "Exento": l["exento"] or None, "Detalle": l["detalle"] or None, "Subcuenta": tr(l["sub"]) or None,
                       "Cuenta manual": tr(l["cue"]) or None, "CC": l["cc"] or None})
for p in pagos:
    P_rows.append({"Fecha": D(p["fecha"]), "Solicitante": tr(p["solicitante"]), "Tipo": p["tipo"], "Banco": tr(p["banco"]), "Monto": p["monto"],
                   "ID rendición": (p["rends"] or [None])[0], "Origen": "Histórico Excel v1.1", "Observación": p["obs"]})
P_rows.sort(key=lambda p: (p["Fecha"] or dt.datetime(2026, 1, 1)))

wb = Workbook()
wsI = wb.active; wsI.title = "Inicio"
def TR(t, c): return f"{t}[[#This Row],[{c}]]"

def build(ws, tname, cols, rows, extra=10):
    """cols: (name, width, kind in/fx/flow, fmt, formula|None)"""
    n = max(len(rows), 1) + (0 if rows else 0)
    for j, (name, w, kind, fmt, fx) in enumerate(cols, 1):
        c = ws.cell(1, j, name); c.font = WH; c.fill = {"in": H_IN, "fx": H_FX, "flow": H_FL}[kind]
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[1].height = 32
    nrows = max(len(rows), 1)
    for i in range(nrows):
        r = i + 2; d = rows[i] if i < len(rows) else {}
        for j, (name, w, kind, fmt, fx) in enumerate(cols, 1):
            c = ws.cell(r, j)
            if fx: c.value = "=" + fx; c.font = BASE
            else: c.value = d.get(name); c.font = BLUE
            if fmt: c.number_format = fmt
    t = Table(displayName=tname, ref=f"A1:{L(len(cols))}{nrows + 1}")
    t._initialise_columns()
    for j, (name, w, kind, fmt, fx) in enumerate(cols):
        t.tableColumns[j].name = name
        if fx: t.tableColumns[j].calculatedColumnFormula = TableFormula(attr_text=fx)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = "C2"
    return nrows + 1

def dv(ws, col, last, src):
    v = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=False); ws.add_data_validation(v); v.add(f"{col}2:{col}{last + 2000}")

# ---------------- Rendiciones
T = "tblRendiciones"; r_ = lambda c: TR(T, c)
RC = [
    ("ID", 11, "flow", None, None), ("Nro", 7, "flow", "0", None), ("Solicitante", 22, "flow", None, None), ("Correo", 26, "flow", None, None),
    ("Empresa", 24, "flow", None, None), ("Fondo", 20, "flow", None, None), ("Fecha", 11, "flow", DATE, None), ("Total formulario", 12, "flow", CLP, None),
    ("Total gastos", 12, "fx", CLP, f"SUMIFS(tblDetalle[Total],tblDetalle[ID],{r_('ID')})"),
    ("Cuadre", 11, "fx", None, f'IF({r_("Total formulario")}="","Sin dato",IF(ABS({r_("Total gastos")}-IFERROR({r_("Total formulario")}*1,0))<=1,"OK","Diferencia"))'),
    ("Estado", 12, "in", None, None), ("Aprobado por", 18, "flow", None, None), ("Fecha aprobación", 11, "flow", DATE, None),
    ("Pagado", 12, "fx", CLP, f'IF({r_("Estado")}<>"Aprobada",0,MAX(0,MIN({r_("Total gastos")},'
        f'SUMIFS(tblPagos[Monto],tblPagos[Solicitante],{r_("Solicitante")})+SUMIFS(tblBanco[Cargo],tblBanco[Asignado a],{r_("Solicitante")},tblBanco[Tipo],"<>No es rendición")'
        f'-SUMIFS({T}[Total gastos],{T}[Solicitante],{r_("Solicitante")},{T}[Estado],"Aprobada",{T}[Fecha],"<"&{r_("Fecha")})'
        f'-SUMIFS({T}[Total gastos],{T}[Solicitante],{r_("Solicitante")},{T}[Estado],"Aprobada",{T}[Fecha],{r_("Fecha")},{T}[ID],"<"&{r_("ID")}))))'),
    ("Saldo por pagar", 12, "fx", CLP, f'IF({r_("Estado")}<>"Aprobada",0,ROUND({r_("Total gastos")}-{r_("Pagado")},0))'),
    ("Estado pago", 15, "fx", None, f'IF({r_("Estado")}<>"Aprobada","—",IF({r_("Total gastos")}=0,"Sin monto",IF({r_("Saldo por pagar")}<=1,"Pagada",IF({r_("Pagado")}>0,"Pago parcial","Pendiente de pago"))))'),
    ("Respaldos", 30, "flow", None, None), ("Reporte PDF", 30, "flow", None, None), ("Origen", 16, "flow", None, None),
    ("Fecha envío", 11, "flow", DATE, None), ("Observaciones", 30, "in", None, None), ("Dudas lectura", 40, "flow", None, None),
]
wsR = wb.create_sheet("Rendiciones"); lastR = build(wsR, T, RC, R_rows)
dv(wsR, "K", lastR, "=Listas!$J$3:$J$6"); dv(wsR, "C", lastR, "=Listas!$A$3:$A$60"); dv(wsR, "E", lastR, "=Listas!$L$3:$L$6")
wsR.conditional_formatting.add(f"K2:K{lastR+2000}", FormulaRule(formula=['K2="Por aprobar"'], fill=amber))
wsR.conditional_formatting.add(f"K2:K{lastR+2000}", FormulaRule(formula=['OR(K2="Observada",K2="Rechazada")'], fill=red))
wsR.conditional_formatting.add(f"P2:P{lastR+2000}", FormulaRule(formula=['P2="Pagada"'], fill=green))
wsR.conditional_formatting.add(f"P2:P{lastR+2000}", FormulaRule(formula=['OR(P2="Pendiente de pago",P2="Pago parcial")'], fill=amber))
wsR.conditional_formatting.add(f"J2:J{lastR+2000}", FormulaRule(formula=['J2="Diferencia"'], fill=red))

# ---------------- Detalle
T2 = "tblDetalle"; d_ = lambda c: TR(T2, c)
num = lambda c: f"IFERROR({d_(c)}*1,0)"
DC = [
    ("ID", 11, "flow", None, None), ("Línea", 6, "flow", "0", None), ("Fecha", 11, "flow", DATE, None), ("Tipo doc", 14, "flow", None, None),
    ("Proveedor", 22, "flow", None, None), ("RUT", 12, "flow", None, None), ("Nro doc", 16, "flow", "@", None),
    ("Neto", 11, "flow", CLP, None), ("IVA", 10, "flow", CLP, None), ("Exento", 11, "flow", CLP, None),
    ("Total", 12, "fx", CLP, f"{num('Neto')}+{num('IVA')}+{num('Exento')}"),
    ("Detalle", 26, "flow", None, None), ("Subcuenta", 24, "flow", None, None), ("Cuenta manual", 16, "in", None, None),
    ("Cuenta", 18, "fx", None, f'IF({d_("Cuenta manual")}<>"",{d_("Cuenta manual")},IFERROR(VLOOKUP({d_("Subcuenta")},tblSubcuentas,2,FALSE),""))'),
    ("CC", 8, "flow", None, None),
    ("Solicitante", 20, "fx", None, f'IFERROR(INDEX(tblRendiciones[Solicitante],MATCH({d_("ID")},tblRendiciones[ID],0)),"")'),
    ("Estado", 12, "fx", None, f'IFERROR(INDEX(tblRendiciones[Estado],MATCH({d_("ID")},tblRendiciones[ID],0)),"")'),
    ("Mes", 11, "fx", None, f'IF({d_("Fecha")}="","",CHOOSE(MONTH({d_("Fecha")}),"enero","febrero","marzo","abril","mayo","junio","julio","agosto","septiembre","octubre","noviembre","diciembre"))'),
    ("Semana", 13, "fx", None, f'IF({d_("Fecha")}="","","Del "&TEXT(DAY({d_("Fecha")}-WEEKDAY({d_("Fecha")},3)),"00")&" al "&TEXT(DAY({d_("Fecha")}-WEEKDAY({d_("Fecha")},3)+4),"00"))'),
]
wsD = wb.create_sheet("Detalle"); lastD = build(wsD, T2, DC, D_rows)
dv(wsD, "D", lastD, "=Listas!$F$3:$F$7"); dv(wsD, "M", lastD, "=Listas!$C$3:$C$80"); dv(wsD, "N", lastD, "=Listas!$H$3:$H$20"); dv(wsD, "P", lastD, "=Listas!$I$3:$I$15")

# ---------------- Pagos
T3 = "tblPagos"
PC = [("Fecha", 11, "in", DATE, None), ("Solicitante", 22, "in", None, None), ("Tipo", 18, "in", None, None), ("Banco", 22, "in", None, None),
      ("Monto", 12, "in", CLP, None), ("ID rendición", 12, "in", None, None), ("Origen", 18, "in", None, None), ("Observación", 30, "in", None, None)]
wsP = wb.create_sheet("Pagos"); lastP = build(wsP, T3, PC, P_rows)
dv(wsP, "B", lastP, "=Listas!$A$3:$A$60"); dv(wsP, "C", lastP, "=Listas!$N$3:$N$6"); dv(wsP, "D", lastP, "=Listas!$P$3:$P$10")

# ---------------- Banco
T4 = "tblBanco"; b_ = lambda c: TR(T4, c)
BC = [("Fecha", 11, "in", DATE, None), ("Banco", 22, "in", None, None), ("Descripción", 40, "in", None, None), ("Nro operación", 14, "in", "@", None),
      ("Cargo", 12, "in", CLP, None), ("Abono", 12, "in", CLP, None), ("Asignado a", 22, "in", None, None), ("Tipo", 16, "in", None, None),
      ("Sugerencia", 34, "fx", None,
       f'IF(OR(IFERROR({b_("Cargo")}*1,0)=0,{b_("Asignado a")}<>""),"",IFERROR("Rendición "&INDEX(tblRendiciones[ID],MATCH(ROUND({b_("Cargo")},0),tblRendiciones[Saldo por pagar],0))&" de "&INDEX(tblRendiciones[Solicitante],MATCH(ROUND({b_("Cargo")},0),tblRendiciones[Saldo por pagar],0)),'
       f'IFERROR("Saldo total de "&INDEX(tblPersonas[Nombre],MATCH(ROUND({b_("Cargo")},0),tblPersonas[Saldo],0)),"")))'),
      ("Estado", 13, "fx", None, f'IF(IFERROR({b_("Cargo")}*1,0)=0,"Abono",IF({b_("Tipo")}="No es rendición","Ajeno",IF({b_("Asignado a")}<>"","Conciliado","Sin conciliar")))')]
wsB = wb.create_sheet("Banco"); lastB = build(wsB, T4, BC, [])
dv(wsB, "B", lastB, "=Listas!$P$3:$P$10"); dv(wsB, "G", lastB, "=Listas!$A$3:$A$60"); dv(wsB, "H", lastB, "=Listas!$R$3:$R$5")
wsB.conditional_formatting.add(f"J2:J{lastB+3000}", FormulaRule(formula=['J2="Sin conciliar"'], fill=amber))
wsB.conditional_formatting.add(f"J2:J{lastB+3000}", FormulaRule(formula=['J2="Conciliado"'], fill=green))

# ---------------- Personas
T5 = "tblPersonas"; p_ = lambda c: TR(T5, c)
PeC = [("Nombre", 22, "in", None, None), ("Código", 8, "in", None, None), ("Correo", 30, "in", None, None), ("Función", 14, "in", None, None),
       ("Rendiciones", 10, "fx", "0", f'COUNTIFS(tblRendiciones[Solicitante],{p_("Nombre")},tblRendiciones[Estado],"Aprobada")'),
       ("Último Nro", 9, "fx", "0", f'IF({p_("Rendiciones")}=0,"",_xlfn.MAXIFS(tblRendiciones[Nro],tblRendiciones[Solicitante],{p_("Nombre")}))'),
       ("Nros faltantes", 10, "fx", "0", f'IF({p_("Rendiciones")}=0,"",MAX(0,{p_("Último Nro")}-COUNTIFS(tblRendiciones[Solicitante],{p_("Nombre")},tblRendiciones[Nro],">=1")))'),
       ("Rendido", 13, "fx", CLP, f'SUMIFS(tblRendiciones[Total gastos],tblRendiciones[Solicitante],{p_("Nombre")},tblRendiciones[Estado],"Aprobada")'),
       ("Pagado", 13, "fx", CLP, f'SUMIFS(tblPagos[Monto],tblPagos[Solicitante],{p_("Nombre")})+SUMIFS(tblBanco[Cargo],tblBanco[Asignado a],{p_("Nombre")},tblBanco[Tipo],"<>No es rendición")'),
       ("Saldo", 13, "fx", CLP, f'ROUND({p_("Rendido")}-{p_("Pagado")},0)'),
       ("Situación", 24, "fx", None, f'IF(AND({p_("Rendido")}=0,{p_("Pagado")}=0),"Sin movimientos",IF({p_("Saldo")}>1,"Empresa le debe",IF({p_("Saldo")}<-1,"Pendiente por rendir","Al día")))'),
       ("Pendientes de pago", 10, "fx", "0", f'COUNTIFS(tblRendiciones[Solicitante],{p_("Nombre")},tblRendiciones[Estado pago],"Pendiente de pago")+COUNTIFS(tblRendiciones[Solicitante],{p_("Nombre")},tblRendiciones[Estado pago],"Pago parcial")'),
       ("Por aprobar", 9, "fx", "0", f'COUNTIFS(tblRendiciones[Solicitante],{p_("Nombre")},tblRendiciones[Estado],"Por aprobar")')]
Pe_rows = [{"Nombre": n, "Código": CODES.get(norm(n), ""), "Función": "Solicitante"} for n in SOL]
wsPe = wb.create_sheet("Personas"); lastPe = build(wsPe, T5, PeC, Pe_rows)
dv(wsPe, "D", lastPe, '"Solicitante,Tesorero"')
wsPe.conditional_formatting.add(f"J2:J{lastPe}", FormulaRule(formula=["J2<-1"], fill=red))
wsPe.conditional_formatting.add(f"J2:J{lastPe}", FormulaRule(formula=["J2>1"], fill=amber))

# ---------------- Listas
wsL = wb.create_sheet("Listas")
blocks = [("A", "Solicitante", SOL), ("C", "Subcuenta", SUB), ("F", "Tipo doc", ["Factura", "Boleta", "Boleta de honorarios", "Transferencia", "No informado"]),
          ("H", "Cuenta", CUE), ("I", "C.C.", CC), ("J", "Estado", ["Por aprobar", "Aprobada", "Observada", "Rechazada"]),
          ("L", "Empresa", [e[0] for e in EMP]), ("N", "Tipo pago", ["Pago", "Anticipo", "No genera pago", "Devolución"]),
          ("P", "Banco", BAN), ("R", "Tipo banco", ["Pago", "Anticipo", "No es rendición"])]
wsL["A1"] = "Listas maestras. Agrega valores al final de cada columna."; wsL["A1"].font = Font(name=F, bold=True, color="1F3864")
for col, head, vals in blocks:
    c = wsL[f"{col}2"]; c.value = head; c.font = WH; c.fill = H_IN; wsL.column_dimensions[col].width = 28
    for k, v in enumerate(vals): wsL[f"{col}{3+k}"] = v; wsL[f"{col}{3+k}"].font = BASE
wsL["M2"] = "RUT"; wsL["M2"].font = WH; wsL["M2"].fill = H_IN
for k, e in enumerate(EMP): wsL[f"M{3+k}"] = e[1]
# Subcuenta -> cuenta sugerida table
wsL["T2"] = "Subcuenta"; wsL["U2"] = "Cuenta sugerida"
for c in ("T2", "U2"): wsL[c].font = WH; wsL[c].fill = H_IN
for k, s in enumerate(SUB):
    wsL[f"T{3+k}"] = s; wsL[f"U{3+k}"] = SUB2.get(s, "")
st = Table(displayName="tblSubcuentas", ref=f"T2:U{2+len(SUB)}"); st.tableStyleInfo = TableStyleInfo(name="TableStyleLight1"); wsL.add_table(st)
wsL.column_dimensions["T"].width = 34; wsL.column_dimensions["U"].width = 22

# ---------------- Inicio
I = wsI; I.sheet_view.showGridLines = False
I.column_dimensions["A"].width = 3; I.column_dimensions["B"].width = 44; I.column_dimensions["C"].width = 18; I.column_dimensions["D"].width = 80
I["B2"] = "Rendiciones Heca Chile"; I["B2"].font = Font(name=F, size=18, bold=True, color="1F3864")
I["B3"] = "Libro de rendiciones POE-GLO-FIN-003 · alimentado por Microsoft Forms y Power Automate"; I["B3"].font = Font(name=F, italic=True, color="595959")
kp = [("Rendiciones por aprobar", '=COUNTIF(tblRendiciones[Estado],"Por aprobar")', "0"),
      ("Rendiciones aprobadas pendientes de pago", '=COUNTIF(tblRendiciones[Estado pago],"Pendiente de pago")+COUNTIF(tblRendiciones[Estado pago],"Pago parcial")', "0"),
      ("Monto por pagar", "=SUM(tblRendiciones[Saldo por pagar])", CLP),
      ("Empresa debe a personas (saldo +)", '=SUMIF(tblPersonas[Saldo],">0")', CLP),
      ("Fondos pendientes por rendir (saldo −)", '=-SUMIF(tblPersonas[Saldo],"<0")', CLP),
      ("Cargos de banco sin conciliar", '=COUNTIF(tblBanco[Estado],"Sin conciliar")', "0"),
      ("Rendiciones que no cuadran con su formulario", '=COUNTIF(tblRendiciones[Cuadre],"Diferencia")', "0"),
      ("Total rendido aprobado", '=SUMIFS(tblRendiciones[Total gastos],tblRendiciones[Estado],"Aprobada")', CLP)]
I["B5"] = "Indicadores"; I["B5"].font = Font(name=F, bold=True, size=12, color="1F3864")
for k, (a, f, fm) in enumerate(kp):
    I.cell(6 + k, 2, a).font = BASE; c = I.cell(6 + k, 3, f); c.number_format = fm; c.font = Font(name=F, bold=True, size=11)
r0 = 16
I.cell(r0, 2, "Cómo se usa").font = Font(name=F, bold=True, size=12, color="1F3864")
steps = [("Persona", "Llena el formulario de Microsoft Forms con su cuenta de Heca: Nº de rendición y fotos del formulario y de sus boletas."),
         ("Lectura automática", "Power Automate lee las fotos con IA y agrega 1 fila en Rendiciones y 1 fila por documento en Detalle, con Estado = Por aprobar."),
         ("Tesorero", "Recibe la aprobación en Teams u Outlook. Si hay que corregir algo, lo corrige aquí en el Excel y luego aprueba o rechaza."),
         ("Reporte", "Al aprobar, el flujo genera el PDF de la rendición con sus respaldos y deja el link en la columna Reporte PDF."),
         ("Banco", "Pega la cartola en la hoja Banco. La columna Sugerencia propone a quién corresponde cada cargo; eliges la persona en 'Asignado a'."),
         ("Personas", "Muestra por persona lo rendido, lo pagado, el saldo y si está al día, por pagar o pendiente por rendir.")]
for k, (a, b) in enumerate(steps):
    I.cell(r0 + 1 + k, 2, a).font = Font(name=F, bold=True); c = I.cell(r0 + 1 + k, 4, b); c.font = BASE; c.alignment = Alignment(wrap_text=True, vertical="top")
    I.row_dimensions[r0 + 1 + k].height = 30
r1 = r0 + 9
I.cell(r1, 2, "Convenciones").font = Font(name=F, bold=True, size=12, color="1F3864")
conv = [("Encabezado celeste", "Lo escribe Power Automate. Se puede corregir a mano antes de aprobar."),
        ("Encabezado azul oscuro", "Lo completa tesorería (Estado, Observaciones, Cuenta manual, pagos y banco)."),
        ("Encabezado gris", "Fórmula. No escribir encima."),
        ("Estado de pago", "Los pagos de cada persona se aplican a sus rendiciones aprobadas desde la más antigua (FIFO). El saldo por persona está en la hoja Personas."),
        ("Historial", f"Incluye {len(R_rows)} rendiciones, {len(D_rows)} líneas y {len(P_rows)} pagos 2026 traídos de Rendiciones_Chile_2026 v1.1.")]
fills = [H_FL, H_IN, H_FX]
for k, (a, b) in enumerate(conv):
    c = I.cell(r1 + 1 + k, 2, a); c.font = WH if k < 3 else Font(name=F, bold=True)
    if k < 3: c.fill = fills[k]
    d = I.cell(r1 + 1 + k, 4, b); d.font = BASE; d.alignment = Alignment(wrap_text=True, vertical="top"); I.row_dimensions[r1 + 1 + k].height = 30

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F: c.font = Font(name=F, size=c.font.size, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.save(OUT)
print(len(R_rows), len(D_rows), len(P_rows))
